import os
import pandas as pd
from docx import Document

def export_results(extracted_data, output_dir, base_filename):
    csv_path = os.path.join(output_dir, f"{base_filename}.csv")
    docx_path = os.path.join(output_dir, f"{base_filename}.docx")
    
    if not extracted_data:
        df = pd.DataFrame(columns=['Block_ID', 'Extracted_Text', 'Confidence'])
        df.to_csv(csv_path, index=False, encoding='utf-8-sig')
    else:
        df = pd.DataFrame(extracted_data)
        df_csv = df[['Block_ID', 'Extracted_Text', 'Confidence']]
        df_csv.to_csv(csv_path, index=False, encoding='utf-8-sig')
        
    doc = Document()
    if extracted_data:
        prev_y = extracted_data[0]['y_min']
        current_paragraph = doc.add_paragraph()
        
        for i, item in enumerate(extracted_data):
            text = item['Extracted_Text']
            current_y = item['y_min']
            
            if i > 0 and current_y - prev_y > 20: 
                current_paragraph = doc.add_paragraph()
            
            current_paragraph.add_run(text + " ")
            prev_y = current_y
            
    doc.save(docx_path)